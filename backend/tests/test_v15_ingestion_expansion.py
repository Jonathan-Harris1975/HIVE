from __future__ import annotations

import base64
import io
import zipfile

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app


def _reset_settings(monkeypatch, tmp_path):
    from app.core.config import get_settings

    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("DATABASE_ENABLED", "true")
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'hive.sqlite3'}")
    monkeypatch.setenv("APP_ENV", "test")
    get_settings.cache_clear()


def test_health_reports_v15_free_tier_limits(monkeypatch, tmp_path) -> None:
    _reset_settings(monkeypatch, tmp_path)
    client = TestClient(app)
    body = client.get("/health").json()
    assert body["build"] == "v1.26.9-review-state-sync"
    assert body["free_tier"]["enabled"] is True
    assert "zip_extract_max_members" in body["free_tier"]["ingestion_limits"]


def test_docx_chunking_uses_document_extractor(monkeypatch, tmp_path) -> None:
    _reset_settings(monkeypatch, tmp_path)
    client = TestClient(app)
    assert client.post("/v1/db/init").json()["ok"] is True

    docx = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("HIVE DOCX report says Vectorize retrieval should use SQL chunks as source of truth.")
    document.save(docx)

    upload = client.post(
        "/v1/files/upload-base64",
        json={
            "filename": "report.docx",
            "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "content_base64": base64.b64encode(docx.read_bytes()).decode("ascii"),
            "test_run_id": "v15-docx",
        },
    )
    assert upload.status_code == 200
    key = upload.json()["file"]["object_key"]

    chunk = client.post("/v1/files/chunk", json={"object_key": key, "test_run_id": "v15-docx"})
    assert chunk.status_code == 200
    body = chunk.json()
    assert body["ok"] is True
    assert body["chunking"]["chunk_count"] >= 1
    preview = body["chunks_preview"][0]["content"]
    assert "Vectorize retrieval" in preview
    assert body["source"]["extraction"]["extractor"] == "docx"


def test_zip_extract_text_chunks_nested_archive(monkeypatch, tmp_path) -> None:
    _reset_settings(monkeypatch, tmp_path)
    client = TestClient(app)
    assert client.post("/v1/db/init").json()["ok"] is True

    nested = io.BytesIO()
    with zipfile.ZipFile(nested, "w") as archive:
        archive.writestr("nested/report.txt", "Nested ZIP says deployment diagnostics need structured logs.")

    outer = io.BytesIO()
    with zipfile.ZipFile(outer, "w") as archive:
        archive.writestr("README.md", "Top-level archive says SQL fallback remains available.")
        archive.writestr("inner.zip", nested.getvalue())

    upload = client.post(
        "/v1/files/upload-base64",
        json={
            "filename": "audit-pack.zip",
            "content_type": "application/zip",
            "content_base64": base64.b64encode(outer.getvalue()).decode("ascii"),
            "test_run_id": "v15-zip",
        },
    )
    assert upload.status_code == 200
    key = upload.json()["file"]["object_key"]

    extract = client.post(
        "/v1/files/zip/extract-text",
        json={
            "object_key": key,
            "recursive": True,
            "chunk": True,
            "max_members": 10,
            "test_run_id": "v15-zip",
        },
    )
    assert extract.status_code == 200
    body = extract.json()
    assert body["ok"] is True
    assert body["extraction"]["summary"]["extracted_count"] == 2
    assert body["extraction"]["summary"]["nested_archives"] == 1
    assert "deployment diagnostics" in body["extracted_text_preview"]
    assert body["chunking"]["ok"] is True

    extracted_key = body["extracted_file"]["object_key"]
    search = client.get(
        "/v1/files/chunks/search",
        params={"key": extracted_key, "query": "deployment diagnostics structured logs", "limit": 3},
    )
    assert search.status_code == 200
    assert search.json()["count"] >= 1


def test_xlsx_extraction_is_bounded(tmp_path) -> None:
    from app.ingestion.text_extractors import extract_text_with_metadata

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Audit"
    for index in range(20):
        sheet.append(["row", index, "Vectorize semantic retrieval"])
    path = tmp_path / "audit.xlsx"
    workbook.save(path)

    result = extract_text_with_metadata(path, xlsx_max_rows_per_sheet=5, xlsx_max_sheets=1)
    assert result.supported is True
    assert result.extractor == "xlsx"
    assert result.truncated is True
    assert "Vectorize semantic retrieval" in result.text
