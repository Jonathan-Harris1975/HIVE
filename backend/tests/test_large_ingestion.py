from __future__ import annotations

import base64
import io
import zipfile

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app


def test_large_text_upload_chunks_successfully(monkeypatch, tmp_path) -> None:
    from app.core.config import get_settings

    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("DATABASE_ENABLED", "true")
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'hive.sqlite3'}")
    get_settings.cache_clear()
    client = TestClient(app)
    assert client.post("/v1/db/init").json()["ok"] is True
    content = "\n\n".join(f"Large document paragraph {i}: deployment diagnostics and Vectorize fallback." for i in range(350))

    upload = client.post(
        "/v1/files/upload-text",
        json={"filename": "large-doc.txt", "content": content, "test_run_id": "large-ingestion-test"},
    )
    assert upload.status_code == 200
    key = upload.json()["file"]["object_key"]

    chunk = client.post(
        "/v1/files/chunk",
        json={"object_key": key, "max_chars": 1200, "overlap_chars": 120, "test_run_id": "large-ingestion-test"},
    )
    assert chunk.status_code == 200
    body = chunk.json()
    assert body["ok"] is True
    assert body["chunking"]["chunk_count"] > 1


def test_larger_zip_inspection_keeps_preview_bounded(monkeypatch, tmp_path) -> None:
    monkeypatch.chdir(tmp_path)
    client = TestClient(app)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for index in range(30):
            archive.writestr(f"docs/item-{index:02d}.txt", "HIVE ZIP ingestion smoke text " * 20)

    upload = client.post(
        "/v1/files/upload-base64",
        json={
            "filename": "larger-sample.zip",
            "content_type": "application/zip",
            "content_base64": base64.b64encode(buffer.getvalue()).decode("ascii"),
            "test_run_id": "larger-zip-test",
        },
    )
    assert upload.status_code == 200
    body = upload.json()
    assert body["file"]["zip_member_count"] == 30
    key = body["file"]["object_key"]

    inspect = client.get("/v1/files/zip/inspect", params={"key": key, "max_members": 10})
    assert inspect.status_code == 200
    zip_body = inspect.json()["zip"]
    assert zip_body["member_count"] == 30
    assert len(zip_body["members_preview"]) == 10


def test_docx_and_xlsx_text_extraction_for_document_ingestion(tmp_path) -> None:
    from app.ingestion.text_extractors import extract_text

    docx_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("HIVE DOCX ingestion paragraph about Vectorize retrieval.")
    document.save(docx_path)

    xlsx_path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Audit"
    sheet.append(["topic", "status"])
    sheet.append(["Vectorize retrieval", "ready"])
    workbook.save(xlsx_path)

    assert "Vectorize retrieval" in extract_text(docx_path)
    assert "Vectorize retrieval" in extract_text(xlsx_path)
